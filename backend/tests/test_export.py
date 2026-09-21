"""GET /annotations/{id}/export over the real request path.

A signed invoice goes through the real pipeline (Step Zero, line items
included), is confirmed through the API, and is exported in all three formats.
Parity is asserted against the JSON contract, the refusal is asserted in both
languages, and every export is checked for its audit row.
"""

from __future__ import annotations

import io
import json
import uuid
from collections.abc import AsyncIterator
from dataclasses import dataclass

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader, PdfWriter
from sqlalchemy import text as sql

from app.api.deps import get_current_org_id
from app.config import Settings, get_settings
from app.db.base import get_sessionmaker
from app.main import create_app
from app.services import pipeline as pipeline_mod
from app.services import storage as storage_mod
from app.services.normalize import has_arabic
from app.services.storage import LocalStorage
from tests import fixtures
from tests.export_parsing import expected_rows, parse_csv, parse_xlsx

pytestmark = pytest.mark.skipif(
    not get_settings().DATABASE_URL,
    reason="DATABASE_URL not set — see backend/.env.example",
)

SCHEMA = {
    "name": "export",
    "fields": [
        {"key": "invoice_number", "type": "string", "label_en": "Invoice number",
         "label_ar": "رقم الفاتورة", "required": True},
        {"key": "seller_name", "type": "string", "label_en": "Seller name",
         "label_ar": "اسم البائع", "required": True},
        {"key": "seller_trn", "type": "string", "label_en": "Seller VAT number",
         "label_ar": "الرقم الضريبي للبائع", "required": True},
        {"key": "subtotal", "type": "decimal", "label_en": "Subtotal",
         "label_ar": "المجموع قبل الضريبة", "required": True},
        {"key": "vat_amount", "type": "decimal", "label_en": "VAT",
         "label_ar": "ضريبة القيمة المضافة", "required": True},
        {"key": "total_amount", "type": "decimal", "label_en": "Total",
         "label_ar": "الإجمالي", "required": True},
    ],
    "line_item_fields": [
        {"key": "line_description", "type": "string", "label_en": "Description",
         "label_ar": "الوصف"},
        {"key": "line_quantity", "type": "decimal", "label_en": "Quantity",
         "label_ar": "الكمية"},
        {"key": "line_amount", "type": "decimal", "label_en": "Line amount",
         "label_ar": "مبلغ البند"},
    ],
}  # fmt: skip


class NoModel:
    def __init__(self, *args: object, **kwargs: object) -> None:
        pass

    def chat(self, **kwargs: object) -> None:
        raise AssertionError("the model must not be called when UBL is present")


@dataclass
class Setup:
    org_id: uuid.UUID
    unconfirmed: uuid.UUID
    confirmed: uuid.UUID
    lifecycle: uuid.UUID
    """Confirmed and never exported: for the status transition."""
    batch_a: uuid.UUID
    batch_b: uuid.UUID
    """Two more, for the multi-invoice workbook."""


def client_as(org_id: uuid.UUID) -> AsyncClient:
    app = create_app(Settings(DEBUG_ENDPOINTS=False))
    app.dependency_overrides[get_current_org_id] = lambda: org_id
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


def _distinct_copy(pdf: bytes, tag: str) -> bytes:
    """Same visible invoice, different bytes: documents has UNIQUE(org_id, sha256)."""
    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(pdf)).pages:
        writer.add_page(page)
    for name, data in (("invoice.xml", fixtures.build_ubl_xml()),):
        writer.add_attachment(name, data)
    writer.add_metadata({"/Subject": tag})
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


async def _process(org_id: uuid.UUID, queue_id: uuid.UUID, pdf: bytes) -> uuid.UUID:
    document_id = uuid.uuid4()
    key = storage_mod.document_key(org_id, document_id, "original.pdf")
    storage_mod.get_storage().put(key, pdf)
    async with get_sessionmaker()() as s, s.begin():
        await s.execute(
            sql(
                "INSERT INTO documents (id, org_id, queue_id, storage_key, source) "
                "VALUES (:i, :o, :q, :k, 'upload')"
            ),
            {"i": document_id, "o": org_id, "q": queue_id, "k": key},
        )
    outcome = await pipeline_mod.process_document(org_id, document_id)
    assert outcome.error is None, outcome.error
    assert outcome.annotation_id is not None
    return outcome.annotation_id


@pytest_asyncio.fixture(scope="module")
async def setup(tmp_path_factory: pytest.TempPathFactory) -> AsyncIterator[Setup]:
    org_id, queue_id = uuid.uuid4(), uuid.uuid4()
    sessionmaker = get_sessionmaker()
    async with sessionmaker() as s, s.begin():
        await s.execute(
            sql("INSERT INTO organizations (id, name) VALUES (:i, :n)"),
            {"i": org_id, "n": f"Export {org_id}"},
        )
        await s.execute(
            sql("INSERT INTO queues (id, org_id, name) VALUES (:i, :o, 'AP')"),
            {"i": queue_id, "o": org_id},
        )
        await s.execute(
            sql(
                "INSERT INTO extraction_schemas (org_id, queue_id, version, definition) "
                "VALUES (:o, :q, 1, CAST(:d AS jsonb))"
            ),
            {"o": org_id, "q": queue_id, "d": json.dumps(SCHEMA, ensure_ascii=False)},
        )

    previous = storage_mod._storage
    storage_mod._storage = LocalStorage(tmp_path_factory.mktemp("export-storage"))
    patch = pytest.MonkeyPatch()
    patch.setattr(pipeline_mod, "OllamaClient", NoModel)
    try:
        base = fixtures.build_pdf_with_text_layer_and_ubl()
        unconfirmed = await _process(org_id, queue_id, _distinct_copy(base, "export-a"))
        confirmed = await _process(org_id, queue_id, _distinct_copy(base, "export-b"))
        lifecycle = await _process(org_id, queue_id, _distinct_copy(base, "export-c"))
        batch_a = await _process(org_id, queue_id, _distinct_copy(base, "export-d"))
        batch_b = await _process(org_id, queue_id, _distinct_copy(base, "export-e"))
        async with client_as(org_id) as client:
            for annotation in (confirmed, lifecycle, batch_a, batch_b):
                response = await client.post(f"/api/v1/annotations/{annotation}/confirm")
                assert response.status_code == 200, response.text
        yield Setup(
            org_id=org_id,
            unconfirmed=unconfirmed,
            confirmed=confirmed,
            lifecycle=lifecycle,
            batch_a=batch_a,
            batch_b=batch_b,
        )
    finally:
        patch.undo()
        storage_mod._storage = previous
        async with sessionmaker() as s, s.begin():
            await s.execute(sql("DELETE FROM organizations WHERE id = :i"), {"i": org_id})


async def _export_rows(annotation_id: uuid.UUID) -> list[tuple[str, str, dict]]:  # type: ignore[type-arg]
    async with get_sessionmaker()() as s:
        rows = (
            await s.execute(
                sql(
                    "SELECT target, status, response FROM exports "
                    "WHERE annotation_id = :a ORDER BY created_at"
                ),
                {"a": annotation_id},
            )
        ).all()
    return [(r.target, r.status, r.response) for r in rows]


# --------------------------------------------------------------------------- #
# Refusal
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("export_format", ["json", "csv", "xlsx"])
async def test_unconfirmed_export_is_refused_in_both_languages(
    setup: Setup, export_format: str
) -> None:
    async with client_as(setup.org_id) as client:
        response = await client.get(
            f"/api/v1/annotations/{setup.unconfirmed}/export", params={"format": export_format}
        )

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["status"] == "to_review"
    assert "confirmed" in detail["message_en"]
    assert has_arabic(detail["message_ar"])
    assert "بانتظار المراجعة" in detail["message_ar"]  # the status, in Arabic too
    assert await _export_rows(setup.unconfirmed) == [], "a refused export was recorded"


async def test_unknown_format_is_rejected(setup: Setup) -> None:
    async with client_as(setup.org_id) as client:
        response = await client.get(
            f"/api/v1/annotations/{setup.confirmed}/export", params={"format": "pdf"}
        )
    assert response.status_code == 422


async def test_export_is_tenant_scoped(setup: Setup) -> None:
    async with client_as(uuid.uuid4()) as client:
        response = await client.get(f"/api/v1/annotations/{setup.confirmed}/export")
    assert response.status_code == 404


# --------------------------------------------------------------------------- #
# Parity, Arabic, audit
# --------------------------------------------------------------------------- #
async def test_all_formats_agree_and_every_export_is_recorded(setup: Setup) -> None:
    bodies: dict[str, bytes] = {}
    async with client_as(setup.org_id) as client:
        for export_format in ("json", "csv", "xlsx"):
            response = await client.get(
                f"/api/v1/annotations/{setup.confirmed}/export",
                params={"format": export_format},
            )
            assert response.status_code == 200, response.text
            assert "attachment" in response.headers["content-disposition"]
            bodies[export_format] = response.content

    expected = expected_rows(bodies["json"])
    assert parse_csv(bodies["csv"]) == expected
    assert parse_xlsx(bodies["xlsx"]) == expected

    contract = json.loads(bodies["json"])
    assert contract["schema_version"] == "munsiq.invoice.v1"
    assert contract["status"] == "confirmed"
    assert [f["key"] for f in contract["invoice"]] == [f["key"] for f in SCHEMA["fields"]]
    for entry in contract["invoice"]:
        assert entry["source"] == "ubl_xml"
        assert entry["original_value"] == entry["value"]

    # Line items came from the signed XML, row by row, in schema order.
    lines = contract["line_items"]
    assert [item["row_index"] for item in lines] == [0, 1]
    assert [f["key"] for f in lines[0]["fields"]] == [
        "line_description",
        "line_quantity",
        "line_amount",
    ]

    recorded = await _export_rows(setup.confirmed)
    assert [(target, status) for target, status, _ in recorded] == [
        ("json", "completed"),
        ("csv", "completed"),
        ("xlsx", "completed"),
    ]
    assert all(r["schema_version"] == "munsiq.invoice.v1" for _, _, r in recorded)
    assert [r["bytes"] for _, _, r in recorded] == [len(bodies[f]) for f in ("json", "csv", "xlsx")]


async def test_arabic_survives_the_real_export(setup: Setup) -> None:
    async with client_as(setup.org_id) as client:
        xlsx = await client.get(
            f"/api/v1/annotations/{setup.confirmed}/export", params={"format": "xlsx"}
        )
        csv_ = await client.get(
            f"/api/v1/annotations/{setup.confirmed}/export", params={"format": "csv"}
        )

    for rows in (parse_xlsx(xlsx.content), parse_csv(csv_.content)):
        seller = next(r for r in rows if r["key"] == "seller_name")
        assert seller["value"] == fixtures.SELLER_NAME
        assert seller["label_ar"] == "اسم البائع"
        described = [r["value"] for r in rows if r["key"] == "line_description"]
        assert "صمام كروي 6 انش" in described


# --------------------------------------------------------------------------- #
# Export is a lifecycle state, not just a download
# --------------------------------------------------------------------------- #
async def _status(annotation_id: uuid.UUID) -> str:
    async with get_sessionmaker()() as s:
        return str(
            await s.scalar(
                sql("SELECT status FROM annotations WHERE id = :a"), {"a": annotation_id}
            )
        )


async def test_a_successful_export_moves_the_annotation_to_exported(setup: Setup) -> None:
    assert await _status(setup.lifecycle) == "confirmed"

    async with client_as(setup.org_id) as client:
        first = await client.get(f"/api/v1/annotations/{setup.lifecycle}/export")
        assert first.status_code == 200
        assert await _status(setup.lifecycle) == "exported"

        # Still downloadable in the other formats, and the contract still says
        # confirmed: the invoice is, it has merely also been exported.
        second = await client.get(
            f"/api/v1/annotations/{setup.lifecycle}/export", params={"format": "xlsx"}
        )
    assert second.status_code == 200
    assert json.loads(first.content)["status"] == "confirmed"
    assert await _status(setup.lifecycle) == "exported"


async def test_a_refused_export_leaves_the_status_alone(setup: Setup) -> None:
    async with client_as(setup.org_id) as client:
        response = await client.get(f"/api/v1/annotations/{setup.unconfirmed}/export")

    assert response.status_code == 409
    assert await _status(setup.unconfirmed) == "to_review"


async def test_confirming_an_exported_annotation_does_not_undo_the_export(
    setup: Setup,
) -> None:
    async with client_as(setup.org_id) as client:
        await client.get(f"/api/v1/annotations/{setup.confirmed}/export")
        response = await client.post(f"/api/v1/annotations/{setup.confirmed}/confirm")

    assert response.status_code == 200
    assert await _status(setup.confirmed) == "exported"


# --------------------------------------------------------------------------- #
# One workbook for several invoices
# --------------------------------------------------------------------------- #
def _batch(*ids: uuid.UUID) -> dict[str, list[str]]:
    return {"ids": [str(i) for i in ids]}


async def test_a_batch_workbook_has_a_header_sheet_and_a_line_items_sheet(setup: Setup) -> None:
    from openpyxl import load_workbook

    async with client_as(setup.org_id) as client:
        response = await client.post(
            "/api/v1/annotations/export", json=_batch(setup.batch_a, setup.batch_b)
        )

    assert response.status_code == 200, response.text
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.xlsx"')
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Invoices", "Line Items"]

    header = list(workbook["Invoices"].iter_rows(values_only=True))
    assert header[0][:2] == ("annotation_id", "document_id")
    assert {row[0] for row in header[1:]} == {str(setup.batch_a), str(setup.batch_b)}
    seller = list(header[0]).index("seller_name")
    # Arabic read back by opening the workbook, as a spreadsheet user would.
    assert all(row[seller] == fixtures.SELLER_NAME for row in header[1:])

    subtotal = workbook["Invoices"].cell(row=2, column=list(header[0]).index("subtotal") + 1)
    assert subtotal.value == 45320, "money is a number, not text"
    assert subtotal.number_format == "#,##0.00", "and it is shown as money"

    lines = list(workbook["Line Items"].iter_rows(values_only=True))
    assert lines[0][:2] == ("annotation_id", "row_index")
    assert len(lines) == 1 + 2 * 2, "two invoices x two line items"
    described = [row[list(lines[0]).index("line_description")] for row in lines[1:]]
    assert described.count("صمام كروي 6 انش") == 2

    assert await _status(setup.batch_a) == await _status(setup.batch_b) == "exported"
    for annotation in (setup.batch_a, setup.batch_b):
        recorded = await _export_rows(annotation)
        assert [target for target, _, _ in recorded] == ["xlsx"]
        assert recorded[0][2]["batch_size"] == 2


async def test_one_unconfirmed_invoice_refuses_the_whole_batch(setup: Setup) -> None:
    async with client_as(setup.org_id) as client:
        response = await client.post(
            "/api/v1/annotations/export", json=_batch(setup.confirmed, setup.unconfirmed)
        )

    assert response.status_code == 409
    detail = response.json()["detail"]
    assert detail["not_confirmed"] == [{"id": str(setup.unconfirmed), "status": "to_review"}]
    assert has_arabic(detail["message_ar"])
    assert await _status(setup.unconfirmed) == "to_review"


async def test_a_batch_is_tenant_scoped(setup: Setup) -> None:
    async with client_as(uuid.uuid4()) as client:
        response = await client.post("/api/v1/annotations/export", json=_batch(setup.confirmed))
    assert response.status_code == 404


async def test_a_batch_needs_at_least_one_id_and_a_bounded_number(setup: Setup) -> None:
    async with client_as(setup.org_id) as client:
        empty = await client.post("/api/v1/annotations/export", json={"ids": []})
        huge = await client.post(
            "/api/v1/annotations/export", json={"ids": [str(uuid.uuid4()) for _ in range(201)]}
        )
    assert empty.status_code == 422
    assert huge.status_code == 422


async def test_a_repeated_id_is_exported_once(setup: Setup) -> None:
    from openpyxl import load_workbook

    async with client_as(setup.org_id) as client:
        response = await client.post(
            "/api/v1/annotations/export", json=_batch(setup.confirmed, setup.confirmed)
        )

    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.content))["Invoices"].iter_rows(values_only=True))
    assert len(rows) == 2, "a header and one invoice"
