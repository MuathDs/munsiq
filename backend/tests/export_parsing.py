"""Read CSV and XLSX exports back into the rows ``flatten`` produces.

Shared by the pure renderer tests and the API tests, so both assert parity with
the same definition of "the same data". These parsers play the part of a
downstream consumer: they know the CSV formula guard and undo it.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import load_workbook

from app.schemas.export import MunsiqInvoiceV1
from app.services.export import (
    COLUMNS,
    INVOICE_SHEET,
    LINE_ITEMS_SHEET,
    Cell,
    flatten,
)

_INTS = {"row_index", "bbox_page"}
_FLOATS = {"bbox_x0", "bbox_y0", "bbox_x1", "bbox_y1"}


def expected_rows(json_body: bytes) -> list[dict[str, Cell]]:
    """The contract, as parsed back from its own JSON."""
    return flatten(MunsiqInvoiceV1.model_validate_json(json_body))


def _typed(column: str, raw: Any) -> Cell:
    if raw is None or raw == "":
        return None
    if column in _INTS:
        return int(raw)
    if column in _FLOATS:
        return float(raw)
    return str(raw)


def parse_csv(body: bytes) -> list[dict[str, Cell]]:
    assert body.startswith(b"\xef\xbb\xbf"), "CSV must carry a UTF-8 BOM for Excel"
    reader = csv.DictReader(io.StringIO(body.decode("utf-8-sig")))
    assert tuple(reader.fieldnames or ()) == COLUMNS
    rows: list[dict[str, Cell]] = []
    for raw in reader:
        row: dict[str, Cell] = {}
        for column in COLUMNS:
            value = raw[column]
            # Undo the formula-injection guard: "'=..." was "=...".
            if value.startswith("'") and value[1:2] in ("=", "+", "-", "@", "\t", "\r"):
                value = value[1:]
            row[column] = _typed(column, value)
        rows.append(row)
    return rows


def parse_xlsx(body: bytes) -> list[dict[str, Cell]]:
    workbook = load_workbook(io.BytesIO(body))
    assert workbook.sheetnames == [INVOICE_SHEET, LINE_ITEMS_SHEET]
    rows: list[dict[str, Cell]] = []
    for sheet_name, section in ((INVOICE_SHEET, "invoice"), (LINE_ITEMS_SHEET, "line_item")):
        sheet = workbook[sheet_name]
        values = list(sheet.iter_rows(values_only=True))
        header = [str(h) for h in values[0]]
        for raw in values[1:]:
            cells = dict(zip(header, raw, strict=True))
            row: dict[str, Cell] = {"section": section, "row_index": None}
            for column in COLUMNS:
                if column == "section" or (column == "row_index" and column not in cells):
                    continue
                row[column] = _typed(column, cells[column])
            rows.append({column: row.get(column) for column in COLUMNS})
    return rows
