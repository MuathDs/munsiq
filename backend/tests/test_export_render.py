"""MunsiqInvoiceV1 renderers — pure functions, no database.

The contract is the model. These tests pin that CSV and XLSX carry exactly what
JSON carries, that Arabic survives each format byte for byte, and that
untrusted invoice text never becomes a spreadsheet formula.
"""

from __future__ import annotations

import io
import json
import uuid
import zipfile
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from app.schemas.export import ExportBBox, ExportField, ExportLineItem, MunsiqInvoiceV1
from app.services.export import (
    INVOICE_SHEET,
    LINE_ITEMS_SHEET,
    neutralise_csv_cell,
    render_csv,
    render_json,
    render_xlsx,
)
from tests.export_parsing import expected_rows, parse_csv, parse_xlsx

SELLER_AR = "شركة الجزيرة للصيانة الصناعية"
LINE_AR = "صمام كروي 6 انش"
FORMULA = '=HYPERLINK("http://example.invalid","x")'
REVIEWER = uuid.UUID("7d0c3c0e-5d7e-4b8e-9a37-2f0f6c1f0a11")


def field(key: str, value: str | None, **overrides: object) -> ExportField:
    base: dict[str, object] = {
        "key": key,
        "label_en": key.replace("_", " ").capitalize(),
        "label_ar": "حقل",
        "type": "string",
        "value": value,
        "source": "ubl_xml",
        "confidence": Decimal("1.0000"),
        "bbox": ExportBBox(page=1, x0=0.1, y0=0.2, x1=0.35, y1=0.225),
        "reviewed_by": None,
        "original_value": value,
    }
    base.update(overrides)
    return ExportField.model_validate(base)


def sample() -> MunsiqInvoiceV1:
    return MunsiqInvoiceV1(
        annotation_id=uuid.UUID("11111111-2222-3333-4444-555555555555"),
        document_id=uuid.UUID("66666666-7777-8888-9999-000000000000"),
        status="confirmed",
        confirmed_at=datetime(2026, 9, 11, 12, 0, tzinfo=UTC),
        confirmed_by=None,
        exported_at=datetime(2026, 9, 11, 12, 5, tzinfo=UTC),
        model_version=None,
        has_embedded_ubl=True,
        invoice=[
            field("invoice_number", "SA-2026-0334", label_ar="رقم الفاتورة"),
            field("seller_name", SELLER_AR, label_ar="اسم البائع", bbox=None),
            field(
                "total_amount",
                "52119.00",
                type="decimal",
                source="human",
                reviewed_by=REVIEWER,
                original_value="52118.00",
            ),
            field(
                "purchase_order_number",
                None,
                source="ocr_rule",
                confidence=Decimal("0.0000"),
                bbox=None,
            ),
            field("buyer_name", FORMULA, source="vlm", confidence=Decimal("0.8123")),
            field("discount", "-150.00", type="decimal"),
        ],
        line_items=[
            ExportLineItem(
                row_index=0,
                fields=[
                    field("line_description", "Centrifugal pump"),
                    field("line_amount", "40000.00", type="decimal"),
                ],
            ),
            ExportLineItem(
                row_index=1,
                fields=[
                    field("line_description", LINE_AR, bbox=None),
                    field("line_amount", "5320.00", type="decimal"),
                ],
            ),
        ],
    )


# --------------------------------------------------------------------------- #
# The contract
# --------------------------------------------------------------------------- #
def test_json_round_trips_and_carries_every_attribute() -> None:
    invoice = sample()
    body = render_json(invoice)
    assert MunsiqInvoiceV1.model_validate_json(body) == invoice

    data = json.loads(body)
    assert data["schema_version"] == "munsiq.invoice.v1"
    required = {"value", "source", "confidence", "bbox", "reviewed_by", "original_value"}
    every = data["invoice"] + [f for item in data["line_items"] for f in item["fields"]]
    for entry in every:
        assert required <= entry.keys(), entry["key"]
    # Arabic is written as UTF-8, not \uXXXX escapes.
    assert SELLER_AR.encode("utf-8") in body


def test_schema_version_is_a_literal() -> None:
    payload = json.loads(render_json(sample()))
    payload["schema_version"] = "munsiq.invoice.v2"
    with pytest.raises(ValidationError):
        MunsiqInvoiceV1.model_validate(payload)


def test_only_confirmed_is_representable() -> None:
    payload = json.loads(render_json(sample()))
    payload["status"] = "to_review"
    with pytest.raises(ValidationError):
        MunsiqInvoiceV1.model_validate(payload)


# --------------------------------------------------------------------------- #
# Parity
# --------------------------------------------------------------------------- #
def test_csv_and_xlsx_carry_exactly_what_json_carries() -> None:
    invoice = sample()
    expected = expected_rows(render_json(invoice))
    assert len(expected) == len(invoice.invoice) + 4

    assert parse_csv(render_csv(invoice)) == expected
    assert parse_xlsx(render_xlsx(invoice)) == expected


def test_money_and_confidence_are_never_floats() -> None:
    workbook = load_workbook(io.BytesIO(render_xlsx(sample())))
    sheet = workbook[INVOICE_SHEET]
    header = [c.value for c in sheet[1]]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        cells = dict(zip(header, row, strict=True))
        assert cells["value"] is None or isinstance(cells["value"], str)
        assert isinstance(cells["confidence"], str)
    totals = [r for r in parse_xlsx(render_xlsx(sample())) if r["key"] == "total_amount"]
    assert totals[0]["value"] == "52119.00"  # not 52119.0


def test_xlsx_has_the_two_sheets_and_the_contract_version() -> None:
    workbook = load_workbook(io.BytesIO(render_xlsx(sample())))
    assert workbook.sheetnames == [INVOICE_SHEET, LINE_ITEMS_SHEET]
    props = {p.name: p.value for p in workbook.custom_doc_props.props}
    assert props["schema_version"] == "munsiq.invoice.v1"
    assert props["annotation_id"] == "11111111-2222-3333-4444-555555555555"


# --------------------------------------------------------------------------- #
# Arabic
# --------------------------------------------------------------------------- #
def test_arabic_round_trips_through_xlsx() -> None:
    body = render_xlsx(sample())
    rows = parse_xlsx(body)
    assert next(r for r in rows if r["key"] == "seller_name")["value"] == SELLER_AR
    lines = [r for r in rows if r["key"] == "line_description"]
    assert [r["value"] for r in lines] == ["Centrifugal pump", LINE_AR]

    # At the byte level too: UTF-8 in the package, not mojibake or escapes.
    with zipfile.ZipFile(io.BytesIO(body)) as package:
        shared = package.read("xl/sharedStrings.xml")
    assert SELLER_AR.encode("utf-8") in shared
    assert LINE_AR.encode("utf-8") in shared

    # And it is laid out right to left.
    sheet = load_workbook(io.BytesIO(body))[INVOICE_SHEET]
    seller = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == "seller_name")
    value_cell = seller[[c.value for c in sheet[1]].index("value")]
    assert value_cell.alignment.readingOrder == 2
    assert value_cell.alignment.horizontal == "right"


def test_arabic_round_trips_through_csv() -> None:
    body = render_csv(sample())
    assert body.startswith(b"\xef\xbb\xbf")
    assert SELLER_AR.encode("utf-8") in body
    rows = parse_csv(body)
    assert next(r for r in rows if r["key"] == "seller_name")["value"] == SELLER_AR


# --------------------------------------------------------------------------- #
# Untrusted text
# --------------------------------------------------------------------------- #
def test_invoice_text_never_becomes_a_formula() -> None:
    sheet = load_workbook(io.BytesIO(render_xlsx(sample())))[INVOICE_SHEET]
    buyer = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == "buyer_name")
    value_cell = buyer[[c.value for c in sheet[1]].index("value")]
    assert value_cell.value == FORMULA
    assert value_cell.data_type == "s", "stored as a formula, not text"

    assert f"'{FORMULA}".replace('"', '""') in render_csv(sample()).decode("utf-8-sig")


@pytest.mark.parametrize(
    ("raw", "written"),
    [
        ("=1+1", "'=1+1"),
        ("+SUM(A1)", "'+SUM(A1)"),
        ("@cmd", "'@cmd"),
        ("-1+1", "'-1+1"),
        ("-150.00", "-150.00"),
        ("-45,320.00", "-45,320.00"),
        ("SA-2026-0334", "SA-2026-0334"),
        ("", ""),
    ],
)
def test_csv_formula_guard(raw: str, written: str) -> None:
    assert neutralise_csv_cell(raw) == written
