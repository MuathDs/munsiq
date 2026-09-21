"""Build MunsiqInvoiceV1 from the database, and render it.

One loader, three renderers. ``load_invoice`` is the only code that reads the
database; ``render_json``, ``render_csv`` and ``render_xlsx`` each take the
finished model and nothing else. CSV and XLSX share ``flatten``, so a column
exists in both or in neither.

Two hazards are handled here rather than left to the consumer, because invoice
text is untrusted input:

* **Formula injection.** A supplier name of ``=HYPERLINK(...)`` must reach the
  spreadsheet as text. XLSX cells are typed as strings explicitly — openpyxl
  would otherwise store any string starting with "=" as a live formula. CSV has
  no types, so a cell that would be read as a formula is prefixed with "'" (the
  OWASP guard); plain negative numbers are left alone.
* **Arabic.** CSV is written as UTF-8 with a BOM, which is what makes Excel
  decode it as UTF-8 instead of the ANSI code page. XLSX is UTF-8 by format;
  Arabic cells are additionally right-aligned with right-to-left reading order.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.export import (
    SCHEMA_VERSION,
    ExportBBox,
    ExportField,
    ExportFormat,
    ExportLineItem,
    MunsiqInvoiceV1,
)
from app.schemas.workspace import bbox_from_json, schema_fields_from_definition
from app.services.normalize import has_arabic

Cell = str | int | float | Decimal | None

COLUMNS: Final[tuple[str, ...]] = (
    "section",
    "row_index",
    "key",
    "label_en",
    "label_ar",
    "type",
    "value",
    "source",
    "confidence",
    "bbox_page",
    "bbox_x0",
    "bbox_y0",
    "bbox_x1",
    "bbox_y1",
    "reviewed_by",
    "original_value",
)
"""CSV columns, in order. XLSX drops the columns a sheet makes redundant."""

INVOICE_SHEET: Final[str] = "Invoice"
LINE_ITEMS_SHEET: Final[str] = "Line Items"
INVOICE_COLUMNS: Final[tuple[str, ...]] = tuple(
    c for c in COLUMNS if c not in ("section", "row_index")
)
LINE_ITEM_COLUMNS: Final[tuple[str, ...]] = tuple(c for c in COLUMNS if c != "section")

_FORMULA_TRIGGERS: Final[str] = "=+-@\t\r"
_PLAIN_NUMBER = re.compile(r"^-?\d[\d,]*(\.\d+)?$")


class ExportNotFoundError(Exception):
    """No such annotation for this tenant (RLS makes the two indistinguishable)."""


class ExportNotConfirmedError(Exception):
    def __init__(self, status: str) -> None:
        super().__init__(status)
        self.status = status


class ExportBatchError(Exception):
    """A batch that cannot be exported as a whole. Nothing is exported."""

    def __init__(
        self, missing: list[uuid.UUID], not_confirmed: list[tuple[uuid.UUID, str]]
    ) -> None:
        super().__init__(f"{len(missing)} missing, {len(not_confirmed)} not confirmed")
        self.missing = missing
        self.not_confirmed = not_confirmed


EXPORTABLE_STATUSES: Final[frozenset[str]] = frozenset({"confirmed", "exported"})
"""An annotation that has already been exported is still confirmed, and may be
downloaded again in another format."""


# --------------------------------------------------------------------------- #
# Load
# --------------------------------------------------------------------------- #
async def load_invoice(
    session: AsyncSession, annotation_id: uuid.UUID, *, now: datetime | None = None
) -> MunsiqInvoiceV1:
    """Assemble the contract for one annotation. Refuses unless confirmed.

    Unfiltered by org_id, like every read in the API: RLS scopes it.
    """
    header = (
        await session.execute(
            text(
                "SELECT a.id, a.document_id, a.status, a.model_version, a.confirmed_at, "
                "a.confirmed_by, a.schema_id, d.has_embedded_ubl "
                "FROM annotations a JOIN documents d ON d.id = a.document_id WHERE a.id = :id"
            ),
            {"id": annotation_id},
        )
    ).first()
    if header is None:
        raise ExportNotFoundError(str(annotation_id))
    if header.status not in EXPORTABLE_STATUSES:
        raise ExportNotConfirmedError(header.status)

    definition = None
    if header.schema_id is not None:
        definition = await session.scalar(
            text("SELECT definition FROM extraction_schemas WHERE id = :s"),
            {"s": header.schema_id},
        )
    specs = {spec.key: spec for spec in schema_fields_from_definition(definition)}
    position = {key: index for index, key in enumerate(specs)}

    rows = (
        await session.execute(
            text(
                "SELECT field_key, row_index, value_extracted, value_final, confidence, "
                "source, bbox FROM extracted_fields WHERE annotation_id = :id"
            ),
            {"id": annotation_id},
        )
    ).all()
    # The most recent correction per field decides who reviewed it.
    reviewers: dict[tuple[str, int | None], uuid.UUID | None] = {
        (r.field_key, r.row_index): r.user_id
        for r in (
            await session.execute(
                text(
                    "SELECT DISTINCT ON (field_key, row_index) field_key, row_index, user_id "
                    "FROM field_corrections WHERE annotation_id = :id "
                    "ORDER BY field_key, row_index, created_at DESC, id DESC"
                ),
                {"id": annotation_id},
            )
        ).all()
    }

    def to_field(row: Any) -> ExportField:
        spec = specs.get(row.field_key)
        key = (row.field_key, row.row_index)
        return ExportField(
            key=row.field_key,
            label_en=spec.label_en if spec else "",
            label_ar=spec.label_ar if spec else "",
            type=spec.type if spec else "string",
            value=row.value_final if row.value_final is not None else row.value_extracted,
            source=row.source,
            confidence=row.confidence,
            bbox=_bbox(row.bbox),
            reviewed_by=reviewers.get(key, header.confirmed_by),
            original_value=row.value_extracted,
        )

    def order(row: Any) -> tuple[int, str]:
        return position.get(row.field_key, len(position)), row.field_key

    header_rows = sorted((r for r in rows if r.row_index is None), key=order)
    line_rows: dict[int, list[Any]] = {}
    for row in rows:
        if row.row_index is not None:
            line_rows.setdefault(row.row_index, []).append(row)

    return MunsiqInvoiceV1(
        annotation_id=header.id,
        document_id=header.document_id,
        status="confirmed",
        confirmed_at=header.confirmed_at,
        confirmed_by=header.confirmed_by,
        exported_at=now or datetime.now(UTC),
        model_version=header.model_version,
        has_embedded_ubl=bool(header.has_embedded_ubl),
        invoice=[to_field(r) for r in header_rows],
        line_items=[
            ExportLineItem(
                row_index=index,
                fields=[to_field(r) for r in sorted(line_rows[index], key=order)],
            )
            for index in sorted(line_rows)
        ],
    )


async def load_batch(
    session: AsyncSession, annotation_ids: list[uuid.UUID], *, now: datetime | None = None
) -> list[MunsiqInvoiceV1]:
    """Load several confirmed annotations, all or none.

    Repeated ids are collapsed (order kept). If ANY id is unknown to this tenant or
    not confirmed, nothing is returned and the error names every offender, so the
    reviewer fixes the selection once rather than one refusal at a time.
    """
    invoices: list[MunsiqInvoiceV1] = []
    missing: list[uuid.UUID] = []
    not_confirmed: list[tuple[uuid.UUID, str]] = []
    for annotation_id in dict.fromkeys(annotation_ids):
        try:
            invoices.append(await load_invoice(session, annotation_id, now=now))
        except ExportNotFoundError:
            missing.append(annotation_id)
        except ExportNotConfirmedError as exc:
            not_confirmed.append((annotation_id, exc.status))
    if missing or not_confirmed:
        raise ExportBatchError(missing, not_confirmed)
    return invoices


async def record_exports(
    session: AsyncSession,
    org_id: uuid.UUID,
    invoices: list[MunsiqInvoiceV1],
    *,
    export_format: ExportFormat,
    body: bytes,
    batch_size: int = 1,
) -> None:
    """Write the audit row for each invoice and move it to 'exported'.

    Export is a lifecycle state, not just a download: it is what tells the next
    reviewer that this invoice has already left the system. Only 'confirmed'
    moves; an annotation that was exported before keeps its status.
    """
    digest = hashlib.sha256(body).hexdigest()
    for invoice in invoices:
        await session.execute(
            text(
                "INSERT INTO exports (org_id, annotation_id, target, status, attempts, response) "
                "VALUES (:org, :a, :target, 'completed', 1, CAST(:response AS jsonb))"
            ),
            {
                "org": org_id,
                "a": invoice.annotation_id,
                "target": export_format,
                "response": json.dumps(
                    {
                        "schema_version": SCHEMA_VERSION,
                        "format": export_format,
                        "bytes": len(body),
                        "sha256": digest,
                        "fields": len(invoice.invoice),
                        "line_items": len(invoice.line_items),
                        "batch_size": batch_size,
                    }
                ),
            },
        )
    await session.execute(
        text(
            "UPDATE annotations SET status = 'exported' "
            "WHERE id = ANY(:ids) AND status = 'confirmed'"
        ),
        {"ids": [i.annotation_id for i in invoices]},
    )


def _bbox(raw: Any) -> ExportBBox | None:
    box = bbox_from_json(raw)
    if box is None:
        return None
    return ExportBBox(page=box.page, x0=box.x0, y0=box.y0, x1=box.x1, y1=box.y1)


# --------------------------------------------------------------------------- #
# Render
# --------------------------------------------------------------------------- #
@dataclass(frozen=True)
class Rendered:
    body: bytes
    media_type: str
    filename: str


_MEDIA_TYPES: Final[dict[str, str]] = {
    "json": "application/json",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def render(invoice: MunsiqInvoiceV1, export_format: ExportFormat) -> Rendered:
    renderer = {"json": render_json, "csv": render_csv, "xlsx": render_xlsx}[export_format]
    return Rendered(
        body=renderer(invoice),
        media_type=_MEDIA_TYPES[export_format],
        filename=f"munsiq-invoice-v1-{invoice.annotation_id}.{export_format}",
    )


def render_json(invoice: MunsiqInvoiceV1) -> bytes:
    # Pydantic writes non-ASCII as-is (no \uXXXX escapes), so Arabic stays legible.
    return invoice.model_dump_json(indent=2).encode("utf-8")


def flatten(invoice: MunsiqInvoiceV1) -> list[dict[str, Cell]]:
    """One row per field: header first, then line items row by row."""
    rows = [_row("invoice", None, f) for f in invoice.invoice]
    for item in invoice.line_items:
        rows.extend(_row("line_item", item.row_index, f) for f in item.fields)
    return rows


def _row(section: str, row_index: int | None, field: ExportField) -> dict[str, Cell]:
    box = field.bbox
    return {
        "section": section,
        "row_index": row_index,
        "key": field.key,
        "label_en": field.label_en,
        "label_ar": field.label_ar,
        "type": field.type,
        "value": field.value,
        "source": field.source,
        # str(Decimal), never float: "1.0000" stays "1.0000".
        "confidence": None if field.confidence is None else str(field.confidence),
        "bbox_page": box.page if box else None,
        "bbox_x0": box.x0 if box else None,
        "bbox_y0": box.y0 if box else None,
        "bbox_x1": box.x1 if box else None,
        "bbox_y1": box.y1 if box else None,
        "reviewed_by": None if field.reviewed_by is None else str(field.reviewed_by),
        "original_value": field.original_value,
    }


def neutralise_csv_cell(value: str) -> str:
    """Prefix a cell a spreadsheet would evaluate as a formula (OWASP CSV injection)."""
    if not value or value[0] not in _FORMULA_TRIGGERS:
        return value
    if value[0] == "-" and _PLAIN_NUMBER.match(value):
        return value  # "-150.00" is a number, not a formula
    return "'" + value


def render_csv(invoice: MunsiqInvoiceV1) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(COLUMNS)
    for row in flatten(invoice):
        writer.writerow(
            [
                ""
                if row[c] is None
                else neutralise_csv_cell(str(row[c]))
                if isinstance(row[c], str)
                else str(row[c])
                for c in COLUMNS
            ]
        )
    # The BOM is what makes Excel read the file as UTF-8 rather than ANSI.
    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(invoice: MunsiqInvoiceV1) -> bytes:
    rows = flatten(invoice)
    workbook = Workbook()
    header_sheet = workbook.active
    if not isinstance(header_sheet, Worksheet):  # pragma: no cover - openpyxl invariant
        raise TypeError("a new workbook always has an active worksheet")
    header_sheet.title = INVOICE_SHEET
    _write_sheet(header_sheet, INVOICE_COLUMNS, [r for r in rows if r["section"] == "invoice"])
    _write_sheet(
        workbook.create_sheet(LINE_ITEMS_SHEET),
        LINE_ITEM_COLUMNS,
        [r for r in rows if r["section"] == "line_item"],
    )

    workbook.properties.title = "MunsiqInvoiceV1"
    workbook.properties.identifier = str(invoice.annotation_id)
    # Not in types-openpyxl's Workbook, but present since openpyxl 3.1.
    custom_props = workbook.custom_doc_props  # type: ignore[attr-defined]
    for name, value in _metadata(invoice).items():
        custom_props.append(StringProperty(name=name, value=value))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


BATCH_INVOICES_SHEET: Final[str] = "Invoices"
BATCH_LEADING_COLUMNS: Final[tuple[str, ...]] = (
    "annotation_id",
    "document_id",
    "confirmed_at",
    "model_version",
    "has_embedded_ubl",
)


def render_batch_xlsx(invoices: list[MunsiqInvoiceV1], *, now: datetime | None = None) -> Rendered:
    """One workbook for several invoices: a row per invoice and a row per line item.

    A different shape from the single export, on purpose. That one is long (a row
    per field, with provenance) because it is a contract for a machine; this is
    wide (a column per field) because it is for a person who wants to sort and sum
    a month of invoices. Values are the reviewed ones. Per-field provenance stays
    in the single export.

    Money is written as a number with two decimals; the Decimal is handed to the
    workbook as-is, never through a float. Everything else is text, typed
    explicitly so that a supplier called "=HYPERLINK(...)" is not a formula.
    """
    stamp = now or datetime.now(UTC)
    header_keys = _ordered_keys(f for inv in invoices for f in inv.invoice)
    line_keys = _ordered_keys(f for inv in invoices for item in inv.line_items for f in item.fields)

    header_rows: list[dict[str, Cell]] = []
    line_rows: list[dict[str, Cell]] = []
    for invoice in invoices:
        row: dict[str, Cell] = {
            "annotation_id": str(invoice.annotation_id),
            "document_id": str(invoice.document_id),
            "confirmed_at": invoice.confirmed_at.isoformat() if invoice.confirmed_at else None,
            "model_version": invoice.model_version,
            "has_embedded_ubl": "true" if invoice.has_embedded_ubl else "false",
        }
        row.update({f.key: _cell(f) for f in invoice.invoice})
        header_rows.append(row)
        for item in invoice.line_items:
            line: dict[str, Cell] = {
                "annotation_id": str(invoice.annotation_id),
                "row_index": item.row_index,
            }
            line.update({f.key: _cell(f) for f in item.fields})
            line_rows.append(line)

    header_columns = (*BATCH_LEADING_COLUMNS, *header_keys)
    line_columns = ("annotation_id", "row_index", *line_keys)
    workbook = Workbook()
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):  # pragma: no cover - openpyxl invariant
        raise TypeError("a new workbook always has an active worksheet")
    sheet.title = BATCH_INVOICES_SHEET
    _write_sheet(
        sheet, header_columns, [{c: r.get(c) for c in header_columns} for r in header_rows]
    )
    _write_sheet(
        workbook.create_sheet(LINE_ITEMS_SHEET),
        line_columns,
        [{c: r.get(c) for c in line_columns} for r in line_rows],
    )

    workbook.properties.title = "MunsiqInvoiceV1 batch"
    custom_props = workbook.custom_doc_props  # type: ignore[attr-defined]
    for name, value in {
        "schema_version": SCHEMA_VERSION,
        "exported_at": stamp.isoformat(),
        "invoice_count": str(len(invoices)),
    }.items():
        custom_props.append(StringProperty(name=name, value=value))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return Rendered(
        body=buffer.getvalue(),
        media_type=_MEDIA_TYPES["xlsx"],
        filename=f"munsiq-invoices-v1-{len(invoices)}-{stamp:%Y%m%d}.xlsx",
    )


def _ordered_keys(fields: Any) -> list[str]:
    """Field keys in the order first seen, which is schema order."""
    return list(dict.fromkeys(f.key for f in fields))


def _cell(field: ExportField) -> Cell | Decimal:
    if field.value is None:
        return None
    if field.type == "decimal":
        try:
            return Decimal(field.value.replace(",", "")).quantize(Decimal("0.01"))
        except InvalidOperation:
            return field.value
    return field.value


def _metadata(invoice: MunsiqInvoiceV1) -> dict[str, str]:
    return {
        "schema_version": SCHEMA_VERSION,
        "annotation_id": str(invoice.annotation_id),
        "document_id": str(invoice.document_id),
        "confirmed_at": invoice.confirmed_at.isoformat() if invoice.confirmed_at else "",
        "exported_at": invoice.exported_at.isoformat(),
        "model_version": invoice.model_version or "",
        "has_embedded_ubl": "true" if invoice.has_embedded_ubl else "false",
    }


_RTL = Alignment(horizontal="right", readingOrder=2)
MONEY_FORMAT: Final[str] = "#,##0.00"


def _write_sheet(sheet: Worksheet, columns: tuple[str, ...], rows: list[dict[str, Cell]]) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row[c] for c in columns])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, Decimal):
                # A real number, shown as money: 18400 would read as a count.
                cell.number_format = MONEY_FORMAT
            elif isinstance(cell.value, str):
                # Always text. openpyxl infers a formula from a leading "=",
                # and invoice text is untrusted.
                cell.data_type = "s"
                if has_arabic(cell.value):
                    cell.alignment = _RTL
    sheet.freeze_panes = "A2"
    for index, column in enumerate(columns, start=1):
        width = 14 if column.startswith("bbox") or column in ("type", "row_index") else 24
        sheet.column_dimensions[get_column_letter(index)].width = width
